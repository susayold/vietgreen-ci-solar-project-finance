from openpyxl.cell.cell import MergedCell
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

# Compatibility shims for the classic workbook builder under openpyxl 3.1.5.
# Merged banner rows return MergedCell objects, and ConditionalFormattingList
# exposes add(rule) rather than an add_color_scale convenience method.
if not hasattr(MergedCell, "column_letter"):
    MergedCell.column_letter = property(lambda self: get_column_letter(self.column))

if not hasattr(ConditionalFormattingList, "add_color_scale"):
    def _add_color_scale(self, sqref, *, start_type, start_color,
                         mid_type=None, mid_value=None, mid_color=None,
                         end_type, end_color):
        rule = ColorScaleRule(
            start_type=start_type,
            start_color=start_color,
            mid_type=mid_type,
            mid_value=mid_value,
            mid_color=mid_color,
            end_type=end_type,
            end_color=end_color,
        )
        self.add(sqref, rule)
    ConditionalFormattingList.add_color_scale = _add_color_scale

import build_classic_project_finance_workbook as model


def build_coverage_returns_fixed(wb):
    """Coverage/returns sheet with the chart helper table outside section rows."""
    ws = wb.create_sheet('08_Coverage_Returns', 8)
    model.banner(
        ws,
        'VIETGREEN PROJECT FINANCE MODEL | COVERAGE & RETURNS',
        'Credit coverage and sponsor returns for the selected project',
        23,
    )
    model.label(ws['A4'], 'Selected Project')
    model.formula(ws['B4'], "='01_Dashboard'!$B$3", align='center', bold=True)

    model.section(ws, 6, 'CREDIT COVERAGE', 23)
    metrics = [
        ('Minimum DSCR', model.project_formula('14_Coverage', 'B'), model.FMT_MULT),
        ('Covenant DSCR', model.project_formula('11_Debt_Terms', 'J', key_col='B'), model.FMT_MULT),
        ('DSCR Headroom', '=B7-B8', model.FMT_MULT),
        ('LLCR', model.project_formula('14_Coverage', 'C'), model.FMT_MULT),
        ('PLCR', model.project_formula('14_Coverage', 'D'), model.FMT_MULT),
        ('DSRA Target (VND bn)', '=' + model.project_formula('14_Coverage', 'E')[1:] + '/1000000000', model.FMT_BN),
    ]
    for r, (name, fx, fmt) in enumerate(metrics, 7):
        model.label(ws.cell(r, 1), name)
        model.formula(ws.cell(r, 2), fx, fmt, bg=model.CYAN, bold=True)
    ws.conditional_formatting.add('B9', model.CellIsRule(operator='lessThan', formula=['0'], fill=model.fill('FFC7CE')))

    model.section(ws, 15, 'PROJECT & EQUITY RETURNS', 23)
    returns = [
        ('Project Cost (VND bn)', "='01_Dashboard'!$E$6", model.FMT_BN),
        ('Equity Required (VND bn)', '=' + model.project_formula('15_Returns_Discount', 'B')[1:] + '/1000000000', model.FMT_BN),
        ('Discount Rate', model.project_formula('15_Returns_Discount', 'C'), model.FMT_PCT2),
        ('Project NPV (VND bn)', '=' + model.project_formula('15_Returns_Discount', 'E')[1:] + '/1000000000', model.FMT_BN),
        ('Equity NPV (VND bn)', '=' + model.project_formula('15_Returns_Discount', 'D')[1:] + '/1000000000', model.FMT_BN),
    ]
    for r, (name, fx, fmt) in enumerate(returns, 16):
        model.label(ws.cell(r, 1), name)
        model.formula(ws.cell(r, 2), fx, fmt, bg=model.YELLOW, bold=True)

    ws['A23'] = 'Year'
    ws['A24'] = 'Project Cash Flow'
    ws['A25'] = 'Equity Cash Flow'
    for r in (23, 24, 25):
        model.style_cell(ws.cell(r, 1), bg=model.PURPLE, color=model.WHITE, bold=True, border=True)
    for i in range(16):
        c = 8 + i
        ws.cell(23, c, i)
        model.style_cell(ws.cell(23, c), bg=model.PURPLE, color=model.WHITE, bold=True, border=True, align='right')
        if i == 0:
            model.formula(ws.cell(24, c), '=-$B$16', model.FMT_BN, bg=model.WHITE, color=model.BLACK)
            model.formula(ws.cell(25, c), '=-$B$17', model.FMT_BN, bg=model.WHITE, color=model.BLACK)
        else:
            pcol = get_column_letter(8 + i)
            wcol = get_column_letter(7 + i)
            model.formula(ws.cell(24, c), f"='05_Project_Workings'!{pcol}20", model.FMT_BN)
            model.formula(ws.cell(25, c), f"='07_Cashflow_Waterfall'!{wcol}19", model.FMT_BN)

    model.label(ws['A28'], 'Project IRR', bg=model.CYAN)
    model.formula(ws['B28'], '=IFERROR(IRR(H24:W24),"n.m.")', model.FMT_PCT2, bg=model.CYAN, color=model.BLACK, bold=True)
    model.label(ws['A29'], 'Equity IRR', bg=model.CYAN)
    model.formula(ws['B29'], '=IFERROR(IRR(H25:W25),"n.m.")', model.FMT_PCT2, bg=model.CYAN, color=model.BLACK, bold=True)

    # Keep chart source helpers below the visible return block so they never
    # collide with the merged section headers at rows 6 and 15.
    helper_row = 33
    ws.cell(helper_row, 4, 'Year')
    ws.cell(helper_row, 5, 'Actual DSCR')
    ws.cell(helper_row, 6, 'Covenant')
    for c in range(4, 7):
        model.style_cell(ws.cell(helper_row, c), bg=model.ORANGE, bold=True, border=True, align='center')
    for i in range(10):
        r = helper_row + 1 + i
        ws.cell(r, 4, i + 1)
        finance_col = get_column_letter(8 + i)
        model.formula(ws.cell(r, 5), f"='06_Financing_Workings'!{finance_col}27", model.FMT_MULT)
        model.formula(ws.cell(r, 6), '=$B$8', model.FMT_MULT)

    chart = model.LineChart()
    chart.title = 'Annual DSCR vs Covenant'
    chart.add_data(model.Reference(ws, min_col=5, max_col=6, min_row=helper_row, max_row=helper_row + 10), titles_from_data=True)
    chart.set_categories(model.Reference(ws, min_col=4, min_row=helper_row + 1, max_row=helper_row + 10))
    chart.height = 7
    chart.width = 14
    ws.add_chart(chart, 'H6')

    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 16
    for c in range(8, 24):
        ws.column_dimensions[get_column_letter(c)].width = 11
    model.configure(ws, 'A4')


model.build_coverage_returns = build_coverage_returns_fixed


if __name__ == '__main__':
    model.main()
